"""Generate the launch-day Word documents from their canonical Markdown files.

This is intentionally a narrow documentation utility. It does not execute,
materialize, or mutate a campaign. The Markdown sources remain canonical.
"""

from __future__ import annotations

import argparse
import os
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


TOOL_ROOT = Path(__file__).resolve().parents[1]
DOCS_ROOT = TOOL_ROOT / "docs" / "automated_evaluation"
OUTPUT_ROOT = DOCS_ROOT / "Word Documents"

# Document skill preset: compact_reference_guide. First-page pattern:
# customer_pack (left-aligned operational title stack). Geometry is explicit so
# Word and LibreOffice produce the same operator artifact.
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
TABLE_CELL_TOP_BOTTOM_DXA = 80
TABLE_CELL_START_END_DXA = 120

DOCUMENTS = {
    "1_launch_control_sheet.docx": "launch_control_sheet.md",
    "2_manual_actions_required.docx": "manual_actions_required.md",
    "3_machine_a_readiness.docx": "machine_a_readiness.md",
    "4_machine_b_setup.docx": "machine_b_setup.md",
    "5_credential_and_asset_setup.docx": "credential_and_asset_setup.md",
    "6_fresh_clone_validation.docx": "fresh_clone_validation.md",
    "7_quick_start_campaign.docx": "quick_start_campaign.md",
    "8_two_machine_launch_runbook.docx": "two_machine_launch_runbook.md",
    "9_massive_campaign_guide.docx": "massive_campaign_guide.md",
    "99_operational_launch_readiness.docx": "operational_launch_readiness.md",
    "999_final_launch_checklist.docx": "final_launch_checklist.md",
}

# Dense operator briefs are allowed deliberate section breaks when an automatic
# split would leave only a small command fragment on the following page.
PAGE_BREAK_BEFORE = {
    "machine_a_readiness.md": {"Final clean-clone evidence"},
}


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-root", type=Path, default=OUTPUT_ROOT)
    return parser


def _set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_margins(
    cell,
    *,
    top_bottom: int = TABLE_CELL_TOP_BOTTOM_DXA,
    start_end: int = TABLE_CELL_START_END_DXA,
) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge in ("top", "start", "bottom", "end"):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        value = top_bottom if edge in {"top", "bottom"} else start_end
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_width(node, tag: str, width: int) -> None:
    element = node.find(qn(tag))
    if element is None:
        element = OxmlElement(tag)
        node.append(element)
    element.set(qn("w:w"), str(width))
    element.set(qn("w:type"), "dxa")


def _table_widths(rows: list[list[str]]) -> list[int]:
    column_count = max(len(row) for row in rows)
    if column_count == 1:
        return [CONTENT_WIDTH_DXA]
    if column_count == 2:
        return [2700, 6660]
    lengths = []
    for column in range(column_count):
        longest = max(
            (
                len(_plain(row[column])) if column < len(row) else 0
                for row in rows
            ),
            default=1,
        )
        lengths.append(max(8, min(longest, 42)))
    widths = [
        max(1080, int(CONTENT_WIDTH_DXA * value / sum(lengths)))
        for value in lengths
    ]
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    if widths[-1] < 1080:
        shortage = 1080 - widths[-1]
        widths[-1] = 1080
        donor = max(range(len(widths) - 1), key=widths.__getitem__)
        widths[donor] -= shortage
    return widths


def _set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    properties = table._tbl.tblPr
    _set_width(properties, "w:tblW", CONTENT_WIDTH_DXA)
    indent = properties.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:w"), str(TABLE_INDENT_DXA))
    indent.set(qn("w:type"), "dxa")
    layout = properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            _set_width(cell._tc.get_or_add_tcPr(), "w:tcW", widths[index])


def _plain(text: str) -> str:
    text = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return text.replace("\\<", "<").replace("\\>", ">")


def _add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.add_run("Page ")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()._r
    run.append(begin)
    run.append(instruction)
    run.append(end)


def _configure_document(document: Document, source_name: str) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in (
        ("Title", 26, "0B2545", 0, 8),
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
    subtitle = document.styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle.font.size = Pt(10)
    subtitle.font.color.rgb = RGBColor.from_string("595959")
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(16)
    subtitle.paragraph_format.line_spacing = 1.0
    for name in ("List Bullet", "List Number"):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header.paragraphs[0]
    header.text = "JUST PEACHY | AUTOMATED SPEECH EVALUATION | CPU + CUDA OPERATIONS"
    header.style = document.styles["Caption"]
    header.runs[0].font.color.rgb = RGBColor(89, 89, 89)
    header.runs[0].font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer_run = footer.add_run(f"Canonical source: {source_name}    ")
    footer_run.font.name = "Calibri"
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor.from_string("595959")
    _add_page_number(footer)
    for run in footer.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string("595959")


def _add_code(document: Document, lines: list[str]) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_geometry(table, [CONTENT_WIDTH_DXA])
    cell = table.cell(0, 0)
    _set_cell_shading(cell, "F2F4F7")
    _set_cell_margins(cell, top_bottom=110, start_end=120)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(7.5)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)


def _add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = document.add_table(rows=0, cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for column, value in enumerate(values):
            cells[column].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cells[column])
            paragraph = cells[column].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(_plain(value.strip()))
            run.font.size = Pt(9)
            if row_index == 0:
                run.bold = True
                run.font.color.rgb = RGBColor.from_string("1F4D78")
                _set_cell_shading(cells[column], "E8EEF5")
        if row_index == 0:
            row_properties = table.rows[0]._tr.get_or_add_trPr()
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            row_properties.append(repeat)
    _set_table_geometry(table, _table_widths(rows))


def _add_blockquote(document: Document, text: str) -> None:
    table = document.add_table(rows=1, cols=1)
    _set_table_geometry(table, [CONTENT_WIDTH_DXA])
    cell = table.cell(0, 0)
    _set_cell_shading(cell, "FFF2CC")
    _set_cell_margins(cell, top_bottom=120, start_end=120)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(_plain(text))
    run.bold = True
    run.font.color.rgb = RGBColor.from_string("7A5A00")


def convert_markdown(source: Path, destination: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    document = Document()
    _configure_document(document, source.name)
    index = 0
    first_heading = True
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            index += 1
            continue
        if stripped.startswith("```"):
            language = stripped[3:].strip()
            code: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code.append(lines[index])
                index += 1
            if language:
                code.insert(0, f"[{language}]")
            _add_code(document, code)
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines):
            separator = lines[index + 1].strip()
            if separator.startswith("|") and set(separator.replace("|", "").replace(":", "").replace("-", "").replace(" ", "")) == set():
                rows = [[cell.strip() for cell in stripped.strip("|").split("|")]]
                index += 2
                while index < len(lines) and lines[index].strip().startswith("|"):
                    rows.append([cell.strip() for cell in lines[index].strip().strip("|").split("|")])
                    index += 1
                _add_table(document, rows)
                continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            level = len(heading.group(1))
            text = _plain(heading.group(2))
            if first_heading and level == 1:
                paragraph = document.add_paragraph(style="Title")
                paragraph.add_run(text)
                subtitle = document.add_paragraph(
                    "Operational contract | generated from version-controlled Markdown"
                )
                subtitle.style = document.styles["Subtitle"]
                first_heading = False
            else:
                if text in PAGE_BREAK_BEFORE.get(source.name, set()):
                    document.add_page_break()
                document.add_heading(text, level=level)
            index += 1
            continue
        if stripped.startswith(">"):
            _add_blockquote(document, stripped.lstrip("> "))
            index += 1
            continue
        checklist = re.match(r"^- \[([ xX])\]\s+(.+)$", stripped)
        if checklist:
            marker = "[x]" if checklist.group(1).lower() == "x" else "[ ]"
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.add_run(f"{marker} {_plain(checklist.group(2))}")
            index += 1
            continue
        bullet = re.match(r"^-\s+(.+)$", stripped)
        if bullet:
            document.add_paragraph(_plain(bullet.group(1)), style="List Bullet")
            index += 1
            continue
        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if numbered:
            document.add_paragraph(_plain(numbered.group(1)), style="List Number")
            index += 1
            continue
        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate or candidate.startswith(("#", "```", "|", ">", "- ")) or re.match(r"^\d+\.\s+", candidate):
                break
            paragraph_lines.append(candidate)
            index += 1
        document.add_paragraph(_plain(" ".join(paragraph_lines)))

    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_name(f".{destination.stem}.tmp.docx")
    document.save(temporary)
    os.replace(temporary, destination)


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    output_root = args.output_root.resolve()
    for output_name, source_name in DOCUMENTS.items():
        source = DOCS_ROOT / source_name
        destination = output_root / output_name
        convert_markdown(source, destination)
        print(f"generated {destination}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
