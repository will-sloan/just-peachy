"""Tests for append-only evidence, idempotence and safe publication."""
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from docx import Document
import report


class ReportTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.repo=Path(self.temp.name)
        self.output=self.repo/'training/whistle_followups_v1/reporting'
        self.output.mkdir(parents=True)

    def notes(self, entries):
        (self.output/'analyst_notes.json').write_text(json.dumps({'entries':entries}))

    def note(self, name):
        return {'id':name,'title':'Observation '+name,'paragraphs':['Evidence for '+name]}

    def test_no_duplicate_after_restart_or_lost_status(self):
        self.notes([self.note('one')])
        self.assertEqual(report.update(self.repo)['added_events'],1)
        target=self.output/report.FILENAME
        original=target.read_bytes()
        (self.output/'report_status.json').unlink()
        self.assertEqual(report.update(self.repo)['added_events'],0)
        self.assertEqual(original,target.read_bytes())

    def test_manual_content_and_previous_findings_survive_append(self):
        self.notes([self.note('one')])
        report.update(self.repo)
        target=self.output/report.FILENAME
        doc=Document(target)
        doc.add_paragraph('Retained human annotation')
        doc.save(target)
        self.notes([self.note('one'),self.note('two')])
        self.assertEqual(report.update(self.repo)['added_events'],1)
        text='\n'.join(p.text for p in Document(target).paragraphs)
        self.assertIn('Retained human annotation',text)
        self.assertEqual(text.count('Evidence for one'),1)
        self.assertEqual(text.count('Evidence for two'),1)

    def test_incomplete_scores_are_not_reported(self):
        initial=self.repo/'training/whistle_age_robust_v1'
        initial.mkdir(parents=True)
        source=initial/'evaluation_summary.json'
        source.write_text('{"complete":false,"metrics":[]}')
        previous=source.read_bytes()
        self.assertEqual(report.collect(self.repo,self.output),[])
        self.assertEqual(source.read_bytes(),previous)

    def test_locked_document_is_retained_and_retried_without_duplicate(self):
        self.notes([self.note('one')])
        report.update(self.repo)
        target=self.output/report.FILENAME
        original=target.read_bytes()
        self.notes([self.note('one'),self.note('two')])
        replace=report.os.replace
        def locked(source,destination):
            if Path(destination)==target: raise PermissionError('Word holds the file')
            return replace(source,destination)
        with patch.object(report.os,'replace',side_effect=locked):
            self.assertEqual(report.update(self.repo)['status'],'waiting_for_document_unlock')
        self.assertEqual(target.read_bytes(),original)
        self.assertTrue((self.output/'Whistle_Fine_Tuning_Findings.pending.docx').exists())
        report.update(self.repo)
        self.assertEqual(report.update(self.repo)['added_events'],0)
        self.assertEqual('\n'.join(p.text for p in Document(target).paragraphs).count('Evidence for two'),1)

    def test_run_path_cannot_escape_campaign(self):
        (self.output.parent/'queue.json').write_text('{"jobs":[{"name":"../../other","spec":{}}]}')
        with self.assertRaises(ValueError): report.collect(self.repo,self.output)

    def test_amendment_and_skip_decision_append_without_claiming_completion(self):
        followup=self.output.parent
        (followup/'extension_snapshot.json').write_text(json.dumps({'plan':{
            'id':'omission_test','rationale':'Marginal contribution is unknown',
            'production_policy':'Keep the first fit unless a new recipe qualifies',
            'comparison_limits':'Limited screening only'}}))
        (followup/'extension_production_decision.json').write_text(json.dumps({
            'rerun_required':False,'selected_production_run':'production_all_audio',
            'reason':'No omission passed all guards','recipe':None}))
        state=report.update(self.repo)
        self.assertEqual(state['added_events'],2)
        self.assertFalse(state['campaign_complete'])
        text='\n'.join(p.text for p in Document(self.output/report.FILENAME).paragraphs)
        self.assertIn('12 training runs',text)
        self.assertIn('optional fit is skipped',text)
        self.assertEqual(report.update(self.repo)['added_events'],0)

    def test_final_table_includes_distinct_first_and_selected_refit(self):
        data={'complete':True,'selected_production_run':'production_selected_omission',
              'productions':{'production_all_audio':{},'production_selected_omission':{}},
              'metrics':[{'model':name,'cohort':'ami','wer':wer} for name,wer in
                         [('whistle_base',.4),('whistle_tuned_best',.35),('whistle_tuned_final',.3),
                          ('whistle_production_all_audio',.28),('whistle_production_selected_omission',.26),
                          ('sherpa_gigaspeech_cached',.45)]]}
        (self.output.parent/'matched_comparison.json').write_text(json.dumps(data))
        report.update(self.repo)
        doc=Document(self.output/report.FILENAME)
        cells=' '.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        self.assertIn('First fit',cells);self.assertIn('Refit',cells)
        self.assertIn('28.0%',cells);self.assertIn('26.0%',cells)
        self.assertIn('refit has lower WER',' '.join(p.text for p in doc.paragraphs))


if __name__=='__main__': unittest.main()
